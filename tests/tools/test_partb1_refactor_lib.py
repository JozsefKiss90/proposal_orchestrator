"""Unit tests for tools/partb1_refactor_lib.py (editing primitives and validation checks)."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from tools import partb1_refactor_lib as lib


# ---------------------------------------------------------------------------
# Markup
# ---------------------------------------------------------------------------
class TestMarkup:
    def test_plain_text_is_one_segment(self):
        assert lib.parse_markup("plain words") == [lib.Segment(text="plain words")]

    def test_bold_and_plain(self):
        segs = lib.parse_markup("**Lead.** rest")
        assert segs == [lib.Segment(text="Lead.", bold=True), lib.Segment(text=" rest")]

    def test_footnote_anchor_attaches_to_preceding_word(self):
        segs = lib.parse_markup("fluctuations [^7]. Next")
        assert segs == [lib.Segment(text="fluctuations"), lib.Segment(footnote=7),
                        lib.Segment(text=". Next")]

    def test_markup_to_plain_strips_everything(self):
        assert lib.markup_to_plain("**A** b[^3] c") == "A b c"


class TestParagraphEditing:
    def test_set_paragraph_markup_builds_runs_and_keeps_style(self):
        doc = Document()
        p = doc.add_paragraph("old text", style="List Bullet")
        lib.set_paragraph_markup(p, "**Bold.** plain[^2] tail")
        assert p.style.name == "List Bullet"
        assert [r.text for r in p.runs if r.text] == ["Bold.", " plain", " tail"]
        assert p.runs[0].bold is True and p.runs[1].bold is None
        refs = p._p.findall(".//" + qn("w:footnoteReference"))
        assert len(refs) == 1 and refs[0].get(qn("w:id")) == "2"
        style = refs[0].getparent().find(qn("w:rPr")).find(qn("w:rStyle"))
        assert style.get(qn("w:val")) == lib.FOOTNOTE_REF_STYLE

    def test_find_paragraph_is_unique_or_raises(self):
        doc = Document()
        doc.add_paragraph("Alpha one")
        doc.add_paragraph("Alpha two")
        doc.add_paragraph("Beta")
        assert lib.find_paragraph(doc, "Beta").text == "Beta"
        with pytest.raises(LookupError):
            lib.find_paragraph(doc, "Alpha")

    def test_find_paragraph_normalises_quotes_and_dashes(self):
        doc = Document()
        doc.add_paragraph("WP1 — Data ’harmonisation’")
        assert lib.find_paragraph(doc, "WP1 - Data 'harmonisation'") is not None

    def test_insert_after_inherits_ppr_and_delete(self):
        doc = Document()
        a = doc.add_paragraph("A", style="List Bullet")
        c = doc.add_paragraph("C")
        b = lib.insert_paragraph_after(a, "B **bold**")
        assert [p.text for p in doc.paragraphs] == ["A", "B bold", "C"]
        assert b.style.name == "List Bullet"
        lib.delete_paragraph(a)
        assert [p.text for p in doc.paragraphs] == ["B bold", "C"]
        assert c.text == "C"


class TestTableEditing:
    @staticmethod
    def _table(doc):
        t = doc.add_table(rows=3, cols=2)
        for i, row in enumerate(t.rows):
            row.cells[0].text = f"R{i}"
            row.cells[1].text = f"content {i}"
        return t

    def test_find_table_and_row(self):
        doc = Document()
        t = self._table(doc)
        assert lib.find_table(doc, "R0") is not None
        assert lib.row_cells(lib.find_row(t, "R2"))[1].text == "content 2"
        with pytest.raises(LookupError):
            lib.find_row(t, "R9")

    def test_set_cell_markup_collapses_to_one_paragraph(self):
        doc = Document()
        t = self._table(doc)
        cell = lib.row_cells(t.rows[1])[1]
        cell.add_paragraph("second para")
        lib.set_cell_markup(cell, "**New** cell")
        assert len(cell.paragraphs) == 1 and cell.text == "New cell"

    def test_set_column_widths_rebalances_grid_and_cells(self):
        doc = Document()
        t = self._table(doc)
        for row in t.rows:
            row.cells[0].width, row.cells[1].width = Cm(6), Cm(12)
        grid = t._tbl.find(qn("w:tblGrid"))
        for col, w in zip(grid.findall(qn("w:gridCol")), (6, 12)):
            col.set(qn("w:w"), str(int(w * 567)))
        lib.set_column_widths(t, [4.0, 14.0])
        assert [round(w, 1) for w in lib.column_widths_cm(t)] == [4.0, 14.0]
        assert round(t.rows[1].cells[1].width.cm, 1) == 14.0
        with pytest.raises(ValueError):
            lib.set_column_widths(t, [4.0, 20.0])

    def test_clone_row_after_and_delete_row(self):
        doc = Document()
        t = self._table(doc)
        t.rows[1]._tr.set("{%s}paraId" % lib.W14_NS, "ABC")
        new = lib.clone_row_after(t, t.rows[1], t.rows[2], ["R3", "content 3"])
        assert [lib.row_cells(r)[0].text for r in t.rows] == ["R0", "R1", "R2", "R3"]
        assert new._tr.get("{%s}paraId" % lib.W14_NS) is None
        with pytest.raises(ValueError):
            lib.clone_row_after(t, t.rows[1], t.rows[2], ["only one"])
        lib.delete_row(t.rows[0])
        assert [lib.row_cells(r)[0].text for r in t.rows] == ["R1", "R2", "R3"]


# ---------------------------------------------------------------------------
# Footnotes and package rewriting
# ---------------------------------------------------------------------------
FOOTNOTES_STUB = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                  b'<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                  b'<w:footnote w:id="1"><w:p><w:r><w:t>one</w:t></w:r></w:p></w:footnote>'
                  b'</w:footnotes>')


class TestFootnotes:
    def test_add_footnotes_appends_in_id_order_with_nine_point_text(self):
        out = lib.add_footnotes(FOOTNOTES_STUB, {8: "Eight & co", 7: "Seven"}).decode()
        assert out.index('w:id="7"') < out.index('w:id="8"') < out.index("</w:footnotes>")
        assert "Eight &amp; co" in out and '<w:sz w:val="18"/>' in out
        assert out.count("<w:footnoteRef/>") == 2

    def test_add_footnotes_rejects_existing_id(self):
        with pytest.raises(ValueError):
            lib.add_footnotes(FOOTNOTES_STUB, {1: "dup"})

    def test_rewrite_docx_parts_transforms_named_part(self, tmp_path: Path):
        doc = Document()
        doc.add_paragraph("x")
        path = tmp_path / "t.docx"
        doc.save(path)
        lib.rewrite_docx_parts(path, {"word/document.xml": lambda b: b.replace(b">x<", b">y<")})
        assert Document(str(path)).paragraphs[0].text == "y"
        with pytest.raises(KeyError):
            lib.rewrite_docx_parts(path, {"word/nope.xml": lambda b: b})


class TestDocxText:
    def test_extracts_paragraphs_and_cells_in_order(self, tmp_path: Path):
        doc = Document()
        doc.add_paragraph("first")
        t = doc.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "c1", "c2"
        doc.add_paragraph("last")
        path = tmp_path / "t.docx"
        doc.save(path)
        text = lib.docx_text(path)
        assert text.paragraphs == ["first", "c1", "c2", "last"]
        assert text.footnotes == {}
        assert "c2\nlast" in text.all_text


# ---------------------------------------------------------------------------
# Text checks
# ---------------------------------------------------------------------------
class TestChecks:
    def test_forbidden_and_required(self):
        text = "The model is frozen at M3 (D1.3). [OWNER-CONFIRM: x]"
        assert lib.check_forbidden(text, [(r"\[OWNER-CONFIRM", "placeholder")]) == [
            "forbidden (placeholder): '\\\\[OWNER-CONFIRM' x1"]
        assert lib.check_forbidden(text, [(r"\bD1\.5\b", "renumbered")]) == []
        assert lib.check_required(text, [("frozen at M3", "freeze")]) == []
        assert lib.check_required(text, [("M4", "freeze")]) == ["missing (freeze): 'M4'"]

    def test_deliverable_ids(self):
        allowed = {"D1.1", "D1.2"}
        assert lib.check_deliverable_ids("D1.1 and D1.2", allowed) == []
        assert lib.check_deliverable_ids("D1.1 and D1.5", allowed) == [
            "unexpected deliverable id D1.5", "deliverable id never mentioned: D1.2"]

    def test_wp_headers(self):
        text = ("WP1 — Data harmonisation | M1–M21 | 5.0 PM.\n"
                "WP4 — AgroVIR placement | M25–M30 | 5.4 PM.")
        assert lib.extract_wp_headers(text) == [("WP1", "Data harmonisation", 1, 21, 5.0),
                                                ("WP4", "AgroVIR placement", 25, 30, 5.4)]

    def test_effort_model_accepts_consistent_matrix(self):
        effort = {"WP1": [2.2, 0, 0.6, 0, 2.2, 0, 0], "WP2": [0.3, 4.2, 0.3, 0.6, 3.6, 0, 0],
                  "WP3": [0.2, 1.2, 1.4, 1.3, 0.1, 1.8, 0], "WP4": [0, 0, 0, 0, 0, 0, 5.4],
                  "WP5": [0.3, 0.6, 0.7, 1.1, 0.1, 1.2, 0.6]}
        totals = {"WP1": 5.0, "WP2": 9.0, "WP3": 6.0, "WP4": 5.4, "WP5": 4.6}
        assert lib.check_effort_model(effort, [3, 6, 3, 3, 6, 3, 6], totals, 30.0) == []

    def test_effort_model_reports_each_mismatch(self):
        effort = {"WP1": [1.0, 1.0], "WP2": [1.0, 1.5]}
        failures = lib.check_effort_model(effort, [2.0, 2.0], {"WP1": 2.0, "WP2": 2.0}, 4.0)
        assert any("WP2" in f for f in failures)
        assert any("period 2" in f for f in failures)
        assert any("grand total" in f for f in failures)

    def test_cut_ledger(self):
        ledger = [lib.Cut("CUT-1", "s1", "gone sentence.", "R07", 1),
                  lib.Cut("CUT-2", "s2", "still here.", "R08", 1),
                  lib.Cut("CUT-2", "s3", "never existed.", "R08", 1)]
        failures = lib.check_cut_ledger(ledger, "gone sentence. still here.", "still here.")
        assert "cut CUT-2 listed more than once" in failures
        assert any("still present" in f for f in failures)
        assert any("not found in baseline" in f for f in failures)
        assert not any("CUT-1" in f for f in failures)

    def test_sentence_lengths_guard_abbreviations(self):
        lengths = lib.sentence_lengths("Prof. Janda supervises. Balanced accuracy is 0.70 here. Done")
        assert [n for n, _ in lengths] == [3, 5, 1]


# ---------------------------------------------------------------------------
# PDF report and Gantt
# ---------------------------------------------------------------------------
class TestPdfReport:
    def test_report_on_generated_pdf(self, tmp_path: Path):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        pdf = tmp_path / "t.pdf"
        c = canvas.Canvas(str(pdf), pagesize=A4)
        for n in (1, 2):
            c.setFont("Helvetica", 11)
            c.drawString(60, 700, f"Body text on page {n}")
            c.setFont("Helvetica", 8.5)
            c.drawString(200, 40, f"Part B - Page {n} of 2")
            c.showPage()
        c.save()
        report = lib.pdf_report(pdf, expected_pages=2, placeholders=["[TBC]"])
        assert report["pages"] == 2 and not report["encrypted"]
        assert [p["footer"] for p in report["per_page"]] == ["Part B - Page 1 of 2",
                                                             "Part B - Page 2 of 2"]
        assert report["per_page"][0]["min_body_pt"] == 11.0
        # base-14 Helvetica is not embedded: the check must say so
        assert any("not embedded" in f for f in report["failures"])
        assert not any("pages" in f for f in report["failures"])
        bad = lib.pdf_report(pdf, expected_pages=3)
        assert "2 pages, expected 3" in bad["failures"]


class TestGantt:
    def test_png_has_requested_pixel_size(self, tmp_path: Path):
        from PIL import Image

        spec = lib.GanttSpec(rows=[("WP1", [(1, 3), (10, 12)], "#1f5474"), ("WP2", [(1, 21)], "#3f7d5c")],
                             milestones=[("1", 3)], deliverables_a=[("1.1", 2)],
                             deliverables_b=[("1.2/3; 5.1", 3)])
        out = tmp_path / "g.png"
        lib.build_gantt_png(out, spec, size_px=(850, 390), dpi=100)
        assert Image.open(out).size == (850, 390)
