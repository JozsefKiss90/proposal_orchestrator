"""Stage-4 condensation builder — FIELDWISE Part B-1 submission form (plan §6).

Renders docs/tier5_deliverables/final_exports/FIELDWISE_Part_B1_manual-condensed_2026-09-06.docx
from the agent-condensed content below. The content is condensed from the resolved run-03
Tier 5 section artifacts (proposal_sections/*.json, post OD round-3 integration, commit
acc5cf3), calibrated against the 2026-08-30 run-02 master for per-section length, table
layout and citation style ONLY (its supervisor/host facts are superseded by the run-03
spine and are NOT carried over).

Display mapping (single authority: workpackage_seed.json _provenance.od1_merge_2026_09_02):
  WP1->WP1, WP3->WP2, WP4->WP3, WP5->WP4; task/deliverable display ids follow the
  displayed WP number. Applied here as DELIVERABLE_DISPLAY; task ids are not used in the
  condensed text at all (described verbally), which removes the merged-WP1 sub-numbering
  ambiguity.

The Gantt figure is generated deterministically from
docs/tier4_orchestration_state/phase_outputs/phase4_gantt_milestones/gantt.json (task
spans aggregated to WP spans) plus the display deliverable/milestone months below.

Self-checks at the end enforce the plan §6.4 pre-freeze conditions that are mechanically
checkable (forbidden phrases, display-id vocabulary, KPI completeness, anchor-markup
absence, word budget). The 10-page print-layout check remains an operator action in Word.

Run: py -3.10 tools/build_partb1_condensed.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
OUT_DOCX = REPO / "docs/tier5_deliverables/final_exports/FIELDWISE_Part_B1_manual-condensed_2026-09-06.docx"
GANTT_PNG = REPO / "docs/tier5_deliverables/final_exports/FIELDWISE_Part_B1_gantt_2026-09-06.png"
GANTT_JSON = REPO / "docs/tier4_orchestration_state/phase_outputs/phase4_gantt_milestones/gantt.json"

# ---------------------------------------------------------------------------
# Display mapping (internal -> evaluator-facing), per od1_merge_2026_09_02
# ---------------------------------------------------------------------------
WP_DISPLAY = {"WP1": "WP1", "WP3": "WP2", "WP4": "WP3", "WP5": "WP4"}
DELIVERABLE_DISPLAY = {
    "D1.1": ("D1.1", 2), "D1.2": ("D1.2", 3), "D1.3": ("D1.3", 3),
    "D2.1": ("D1.4", 3), "D2.2": ("D1.5", 14),
    "D3.1": ("D2.1", 18), "D3.2": ("D2.2", 18),
    "D4.1": ("D3.1", 24), "D4.2": ("D3.2", 24),
    "D5.1": ("D4.1", 30), "D5.2": ("D4.2", 30),
}
MILESTONES = [  # (id, month, displayed WP row)
    ("MS1", 2, "WP1"), ("MS6", 3, "WP1"), ("MS2", 14, "WP1"),
    ("MS3", 18, "WP2"), ("MS4", 24, "WP3"), ("MS5", 30, "WP4"),
]

# ---------------------------------------------------------------------------
# Condensed content. P = paragraph (optional bold lead-in). T = table.
# ---------------------------------------------------------------------------
C = {"blocks": []}


def H(level: int, text: str) -> None:
    C["blocks"].append({"kind": "h", "level": level, "text": text})


def P(text: str, lead: str = "") -> None:
    C["blocks"].append({"kind": "p", "lead": lead, "text": text})


def T(header: list, rows: list, widths: list | None = None) -> None:
    C["blocks"].append({"kind": "t", "header": header, "rows": rows, "widths": widths})


def IMG(path: Path, caption: str) -> None:
    C["blocks"].append({"kind": "img", "path": str(path), "caption": caption})


H(0, "Part B-1")

# ============================ 1. EXCELLENCE ============================
H(1, "1. Excellence")
H(2, "1.1 Quality and pertinence of the project's research and innovation objectives (and the extent to which they are ambitious, and go beyond the state of the art)")

P(lead="Problem and overarching aim.", text=(
    "Irrigation decisions in high-value horticulture are still frequently taken from delayed visual "
    "symptoms, isolated sensor thresholds or retrospective crop-performance records, while water "
    "scarcity and climatic variability raise the cost of every wrong decision. FIELDWISE answers this "
    "with a 30-month MSCA European Postdoctoral Fellowship hosted at HUN-REN Agrartudomanyi "
    "Kutatokozpont (ATK), Martonvásár, Hungary, carried out by Dr. Rositsa Cholakova. The action asks "
    "one central question: can a water-stress signal discovered in a data-rich processing-tomato system "
    "remain biologically meaningful across seasons and sensing scales, survive prospective validation "
    "under real farmer conditions, and be encoded in a crop-configurable decision-support framework? It "
    "draws its evidence from an existing five-season archive rather than a new campaign, and channels "
    "inter-sectoral exposure through an integral end-of-project non-academic placement."))

P(lead="Pertinence.", text=(
    "The objectives are pertinent on both axes the Excellence criterion weighs. Scientifically, they "
    "close three focused gaps: the physiological grounding of water-stress indicators — spectral and "
    "environmental variables are too often treated as stress proxies without demonstrating that they "
    "track an actual plant response; transferability across unseen conditions, scales and crops; and "
    "irrigation decision relevance. In instrument terms, the same objectives carry the fellow's "
    "documented transition from plant physiologist and agronomist toward an independent Agricultural "
    "Data Scientist — precisely the interdisciplinary and inter-sectoral competences the call's "
    "expected outcomes seek — so the scientific and career objectives reinforce one another."))

P(lead="The four objectives, and why they are measurable and verifiable.", text=(
    "Each objective carries a stated, checkable output anchored to a dated deliverable. O1 — Harmonise "
    "the heterogeneous 2022–2026 MATE archive and define a physiologically meaningful water-stress "
    "target: the harmonised database and variable-availability map (D1.1, month 2) and the "
    "stress-target, feature-selection and validation protocol (D1.2, month 3), both before the model "
    "freeze at the end of month 3. O2 — Develop a parsimonious, interpretable and uncertainty-aware "
    "predictive framework under HUN-REN supervision, with ELTE spectral/EO harmonisation, separating "
    "reusable components from crop-specific parameters: the primary model, benchmark and model card "
    "(D1.4, month 3) and the spectral/UAV/Sentinel harmonisation and crop-transfer specification (D1.5, "
    "month 14). O3 — Prospectively validate the framework at an independent commercial farmer site "
    "equipped with project meteorological and soil sensing: the validation dataset (D2.1) and the "
    "prospective transferability, uncertainty and limited-recalibration report (D2.2), month 18. O4 — "
    "Operationalise validated outputs in a crop-configurable DrR architecture through continuous "
    "Krumatic development and AgroVIR requirements feedback, followed by industrial evaluation: the DrR "
    "web MVP and pre-placement dossier (D3.1, D3.2, month 24) and the placement assessment and roadmap "
    "(D4.1, D4.2, month 30)."))

P(text=(
    "Measurability is reinforced by two design choices. First, blocked, leakage-safe unseen-year "
    "validation against a pre-specified physiological stress target, using pre-registered metrics — "
    "balanced accuracy, F1, ROC-AUC (detection); RMSE, MAE, R² (prediction); reliability, Brier score, "
    "interval coverage (calibration); lead time and false-alarm rate (operational value) — with "
    "literature-derived acceptance thresholds fixed in D1.2 before the prospective season. Second, the "
    "primary model is frozen and its model card registered (D1.4, MS6) before the first field season, "
    "converting every objective into a falsifiable, auditable claim."))

P(lead="Realistically achievable.", text=(
    "The objectives are ambitious but bounded within the 30-month action. O1 and O2 rest on the "
    "existing MATE archive, developed under ATK supervision (primary supervisor Prof. Tibor Janda) "
    "with ELTE spectral/EO harmonisation. O3 is a single, focused prospective test — the minimum the "
    "logic requires — protected by a defined site-contingency ladder. O4 builds on the fellow's "
    "pre-existing, self-financed research tool DrR – Digital Agronomist, today a monitoring and "
    "data-integration tool rather than a validated advisory product: only validated components are "
    "incorporated, Krumatic supports software engineering from project start, and AgroVIR is consulted "
    "from month 1. Validation within the action is limited to processing tomato — chosen because "
    "roughly five seasons of data, expertise and collaborations already exist for it — while the "
    "crop-transfer specification (D1.5) defines the gates any later crop must pass; the follow-on "
    "route to Bulgarian conditions sits after the action through MVCRI, with no in-action MVCRI effort "
    "claimed."))

P(lead="Beyond the state of the art, and the extent of the ambition.", text=(
    "The ambition lies in refusing four assumptions routine in the field: that a vegetation index "
    "equals water stress; that proximal and satellite indices are interchangeable; that internal "
    "cross-validation proves field transferability; and that a tomato-trained model transfers "
    "unchanged to another crop. The advance is a chain of transitions: from stress proxy to "
    "physiologically defined target; from internal fit through blocked unseen-year validation to "
    "independent commercial-field validation; from proximal sensing through an ELTE hyperspectral/UAV "
    "bridge to Sentinel-2 field-scale inputs; from a tomato-specific prototype to a crop-configurable "
    "framework. The novelty is not a new machine-learning algorithm — the stance is parsimony and "
    "interpretability, no deep-learning dependency — but a validation framework judging operational "
    "data streams against independent plant-level ground truth under unseen-year and new-field "
    "conditions. Concretely, FIELDWISE yields three citable outputs that do not currently exist for "
    "irrigated processing tomato: a physiologically grounded water-stress prediction model with "
    "documented calibration, uncertainty and unseen-year performance (D1.4); a prospectively "
    "quantified research-to-commercial-field transfer gap, measured against pre-registered criteria "
    "rather than estimated (D2.2); and a crop-transfer specification separating reusable elements from "
    "crop-specific components requiring recalibration (D1.5). The second is what is currently not "
    "achievable: published models for this crop report within-experiment performance; the transfer "
    "loss is asserted, not prospectively measured. The principal new knowledge is quantitative "
    "evidence of whether, and how, a biologically grounded water-stress model remains valid when "
    "transferred from experimental data to operational commercial-field conditions."))

H(2, "1.2 Soundness of the proposed methodology (including interdisciplinary approaches, consideration of the gender dimension and other diversity aspects if relevant for the research project, and the quality of open science practices)")

P(lead="Overall methodology: concepts, models and assumptions.", text=(
    "FIELDWISE rests on one conviction: a water-stress model is only useful for irrigation if judged "
    "against four linked criteria at once — physiological validity, predictive performance, "
    "transferability to unseen conditions and management relevance — rather than retrospective "
    "accuracy alone. Concretely, the model learns the relationship between operationally obtainable "
    "predictors (soil-water, meteorological and spectral/Earth-observation variables) and an "
    "independently measured physiological water-stress target. Four stated working hypotheses "
    "(physiological signature, transferability, scale transition, operational relevance) make the "
    "underlying biology explicit and falsifiable: each can fail visibly."))

P(text=(
    "Stage 1 (WP1) establishes the empirical substrate and the model. The 2022–2026 MATE archive is "
    "harmonised into a QC-controlled common-variable matrix with an explicit missingness map, and a "
    "pre-specified stress-target hierarchy is fixed in the protocol before any model fitting — any "
    "substitution forced by the archive audit is documented before results are inspected, never after. "
    "Front-loading makes this ordering structural: the harmonised knowledge base (D1.1) "
    "lands at month 2 and the validation protocol (D1.2) at month 3, both before the freeze; the "
    "primary, parsimonious, uncertainty-aware model — with a single nonlinear benchmark where justified "
    "— is frozen at the end of month 3 (MS6). No outcome data from the commercial field are used before "
    "the freeze. The scale-harmonisation strand continues to month 14, delivering the crop-transfer "
    "specification (D1.5)."))

P(text=(
    "Stage 2 (WP2) treats transferability at three distinct levels. Temporal transfer is tested by "
    "blocked, leakage-safe unseen-year validation. Environmental transfer is tested by prospectively "
    "applying the frozen model at an independent commercial farmer field instrumented by the project, "
    "supported by physiological, proximal, UAV and Sentinel observations — quantifying the transfer "
    "gap and applying only a limited, pre-specified recalibration if transfer is incomplete (D2.1, "
    "D2.2, month 18). Crop transfer is handled by separating a reusable decision-support core — data "
    "architecture, physiology-to-prediction workflow, uncertainty logic, sensing-scale integration, "
    "irrigation-attention interface — from crop-specific reference ranges, thresholds and calibration "
    "parameters, with explicit recalibration gates (D1.5)."))

P(text=(
    "Stage 3 (WP3, then WP4) operationalises: the pre-existing DrR prototype incrementally absorbs "
    "only validated functions into a crop-configurable web MVP (D3.1), which AgroVIR evaluates for "
    "usability and interoperability during the placement (D4.1, D4.2). The software follows the "
    "science, not the other way around."))

P(lead="Methodological challenges and how they are overcome.", text=(
    "(i) Archive variables differ among years: an explicit year-by-variable matrix, harmonising only "
    "defensible common variables. (ii) Optimistic generalisation from random splits: blocked "
    "unseen-year validation, pre-registered thresholds. (iii) The research-to-field gap: a frozen "
    "model, prospective testing and a pre-specified limited-recalibration protocol, so any performance "
    "drop is measured, not hidden. (iv) Fragile UAV campaigns: acquisitions prioritised around key "
    "physiological dates; physiology and proximal sensing remain the core pathway. (v) A single commercial site: a contingency ladder — a second processing-tomato producer "
    "as claim-preserving backup, a public-garden route as measurement-only fallback with a downgraded "
    "transferability claim. (vi) Over-fitting for accuracy: parsimony, with operational value (lead "
    "time, false-alarm rate) scored alongside statistical fit."))

P(lead="Integration of methods and disciplines.", text=(
    "The central challenge runs from plant biology through statistical modelling and remote sensing to "
    "operational software; no single discipline can carry it. Plant-stress physiology and agronomy "
    "supply the stress target and the judgement of whether signals track genuine plant responses, "
    "anchored in the host's Plant Physiology and Metabolomics Department and the MATE archive base. "
    "Agricultural data science supplies parsimonious modelling, calibration, uncertainty and "
    "leakage-safe validation design — carried on the ELTE side by a named predictive-modelling team "
    "member. Remote sensing and geoinformatics supply the hyperspectral/UAV-to-Sentinel scale "
    "transition; software engineering and farm-management expertise (Krumatic, AgroVIR) translate "
    "validated functions into a maintainable prototype. The disciplines meet at defined interfaces — "
    "the target constrains the predictors; the protocol constrains the model; the scale specification "
    "governs field inputs — with the fellow as the integrating agent."))

P(lead="Gender dimension and other diversity aspects.", text=(
    "A sex/gender dimension in the research content is not relevant: the subject is the "
    "processing-tomato crop and its physiological water-stress response, with no human or animal "
    "biological data on which a sex/gender analysis could operate. Diversity is addressed where the "
    "content touches people: the decision-support output is designed user-centredly for growers with "
    "differing farm sizes, digital literacy and contexts; the uptake pathway is bilingual and "
    "cross-border (Hungarian and Bulgarian grower communities); grower feedback is governed by "
    "informed consent, withdrawal rights and coded or anonymised processing."))

P(lead="Open science practices.", text=(
    "Open science is built into the workflow under 'as open as possible, as closed as necessary'. "
    "Pre-registration of the stress target, metrics and thresholds (D1.2) guards against "
    "outcome-selective reporting; the primary model ships with a model card (D1.4); the harmonised "
    "archive is FAIR-by-design; field-scale inputs use open Copernicus/Sentinel-2 data; the measured "
    "transfer gap is published, not hidden. At least two open-access peer-reviewed papers are targeted, "
    "with a possible third methods paper. A Data Management Plan is produced within the first six "
    "months, detailing data types, standards, storage, sharing and reuse. Data governance separates "
    "four rights layers — DrR background, FIELDWISE foreground, institutional datasets, partner code — "
    "with potentially protectable foreground evaluated before disclosure; results from commercially "
    "sensitive farmer or partner information appear only in aggregated, coded or anonymised form."))

H(2, "1.3 Quality of the supervision, training and of the two-way transfer of knowledge between the researcher and the host")

P(lead="Supervisory architecture and qualifications.", text=(
    "Supervision is structured around the action's two disciplinary pillars. Primary supervision at "
    "the recruiting organisation, HUN-REN ATK, is provided by Prof. Tibor Janda, Head of the Plant "
    "Physiology and Metabolomics Department at the Agricultural Institute, whose plant-physiology, "
    "abiotic-stress (including drought) and stress-metabolomics expertise anchors the stress-target "
    "definition, the physiological-validity judgement and the measurement programme (≈2.5 "
    "person-months). Co-supervision is provided by Prof. András Jung (ELTE, Institute of Cartography "
    "and Geoinformatics) — remote sensing, hyperspectral imaging, field spectroscopy, multisensor "
    "fusion — carrying the scale-transition strand (≈2.0 person-months). ELTE additionally contributes "
    "Dr. Roland Hollós as the named team member for the predictive-modelling/machine-learning strand, "
    "and Dr. Zsófia Varga for UAV operations; MATE contributes the five-season archive and agronomic "
    "expertise through Dr. Sándor Takács, with precision-agriculture consultation from Prof. Gábor "
    "Milics. The host's capacity is evidenced, not asserted: ATK coordinated the MSCA Individual "
    "Fellowship LANDRACES (752453, 2017–2019) and participates in Horizon Europe today (COUSIN, "
    "AI4SoilHealth, TUdi)."))

P(lead="Structured supervision and governance.", text=(
    "A defined cadence replaces ad hoc contact: weekly one-to-one supervision; technical sessions with "
    "the co-supervisor and the ELTE-side modelling team member; a monthly written progress record; "
    "quarterly reviews against the Career Development Plan. A Project Steering Group — supervisor "
    "(chair), fellow, co-supervisor, modelling team member, one named contact per partner — issues "
    "go/no-go verdicts at M6, M14, M18, M24 and M27, including a mid-placement review, with pre-defined "
    "decision rights and a two-step conflict-resolution route — a stable support structure for a "
    "returning researcher that protects her scientific independence."))

P(lead="Planned training activities.", text=(
    "Training is supervised practice on real project work: (i) physiological water-stress "
    "interpretation and phenotyping at the host; (ii) geoinformatics, field spectroscopy, "
    "hyperspectral/UAV methods and scale-aware EO interpretation at ELTE; (iii) heterogeneous-data "
    "integration, AI/ML-supported modelling, leakage-safe validation design and uncertainty analysis "
    "with the ELTE-side modelling team member — together closing the competence gap and completing the "
    "progression to an independent Agricultural Data Scientist. Transferable skills are trained "
    "explicitly: research integrity; open science and FAIR practice; generative-AI literacy; grant "
    "writing; IP and exploitation; project management; science communication. A teaching strand "
    "carries five years of university teaching into the new discipline via guest lectures and "
    "practicals at MATE and, where possible, AU Plovdiv. "
    "The plan is consolidated in the Career Development Plan (D1.3, agreed by month 3, signed by both "
    "supervisors and the fellow) and reviewed at the five steering points."))

P(lead="Two-way transfer of knowledge.", text=(
    "Into the fellow flow advanced agricultural data science, robust interpretable modelling, "
    "uncertainty analysis, hyperspectral/UAV/EO integration and research-to-industry translation. Into "
    "the host and partners flow agronomy, plant protection, plant-stress biology, physiological "
    "phenotyping, field experimentation, proximal spectral interpretation and the DrR concept — the "
    "biological grounding that keeps the modelling and EO strands physiologically valid. This genuine "
    "exchange is the knowledge-transfer and capacity outcome the call seeks for participating "
    "organisations."))

P(lead="Rationale and added value of the non-academic placement.", text=(
    "A six-month placement at AgroVIR (AGROVIR Üzletviteli Tanácsadó Kft., months 25–30), an integral, "
    "separately evaluated part of the proposal, supplies inter-sectoral competences the academic host "
    "cannot provide: operational model evaluation against a production farm-management information "
    "system, requirements engineering, interoperability assessment and exploitation planning. Its "
    "added value is operational: an FMIS reported in production use on more than 745,000 ha in eight "
    "countries including Hungary and Bulgaria (company-declared, pending written confirmation). The placement is delivered through monthly structured face-to-face technical "
    "reviews with continuous remote collaboration, supervised by Miklós Maróti (Managing Director), and "
    "is designed around protected, coded data extracts, so no access to proprietary customer systems is "
    "required. AgroVIR is consulted from month 1, and a written placement agreement concluded "
    "beforehand fixes result ownership (grant-agreement default, vesting in HUN-REN ATK) and DrR "
    "background access."))

P(text=(
    "The status of forward-looking commitments is stated transparently: the supervisors' formal "
    "acceptances and efforts, the ELTE-side modelling representation and instrument access, the "
    "data-partner contact and the placement arrangement are operator-declared working assumptions to "
    "be confirmed in partner writing before the action starts; the identities, roles and "
    "infrastructure they rest on are confirmed on public and partner-authored evidence."))

H(2, "1.4 Quality and appropriateness of the researcher's professional experience, competences and skills")

P(lead="A profile matched to the project's interdisciplinary chain.", text=(
    "Dr. Rositsa Cholakova (publishing before 2020 as Rositsa Cholakova-Bimbalova) is strong precisely "
    "where the science demands it: BSc/MSc in Plant Protection, PhD in Plant Physiology (2020, "
    "Agricultural University of Plovdiv), Assistant Professor in Plant Physiology and Plant Stress "
    "Physiology in Plovdiv 2015–2020, and since 2024 Chief Assistant Professor at the Maritsa "
    "Vegetable Crops Research Institute, with current abiotic-stress work in Solanum lycopersicum. "
    "This is the competence that lets her judge whether a model-derived or remotely sensed signal "
    "tracks a genuine plant response — the discriminating judgement the whole validation chain rests "
    "on."))

P(lead="Data-science and EO competences already in place.", text=(
    "Her transition into the second pillar is already under way, which makes it credible rather than "
    "aspirational: a Data Analyst period at KPMG Global Hungary (Python, R, machine-learning basics); "
    "the self-financed DrR – Digital Agronomist research tool she already uses at MVCRI; a COST PANGEOS "
    "Short-Term Scientific Mission at MATE (2025); Erasmus+ mobilities to MATE and Novi Sad "
    "(2025–2026); an ELTE internship (2026); an A1/A3 UAV pilot licence; an in-progress second MSc in "
    "Environmental Engineering; and membership of three COST Actions. Five years of university "
    "lecturing in Bulgarian and English supply the teaching and communication base."))

P(lead="Honest appraisal of fit.", text=(
    "The publication record is concentrated in plant physiology; the remote-sensing items are recent "
    "and co-authored — the EO half of the profile rests on practical, mobility- and licence-based "
    "experience, and the proposal says so openly. The one element not yet at the required level is "
    "rigorous predictive modelling and validation across time, scales and environments — the principal "
    "methodological competence gap the fellowship closes through supervised practice (section 1.3). "
    "She is a returning researcher: after a documented 34-month full-time-parenting career break "
    "(2021–2024) she has re-entered research with clear momentum, and the break does not count against "
    "the MSCA research-experience window."))

# ============================ 2. IMPACT ============================
H(1, "2. Impact")
H(2, "2.1 Credibility of the measures to enhance the career perspectives and employability of the researcher and contribution to his/her skills development")

P(text=(
    "The measures are designed around one evidence-based objective: closing the principal "
    "methodological competence gap — rigorous predictive modelling and validation across time, sensing "
    "scales and environments — that stands between an accomplished plant physiologist/agronomist and an "
    "independent Agricultural Data Scientist. Two concrete post-fellowship trajectories anchor the "
    "career logic: an independent Agricultural Data Scientist or researcher role at the interface of "
    "plant science, Earth observation and agricultural technology, building on the FIELDWISE modelling "
    "record and the AgroVIR placement; and an independent research and innovation line in Bulgaria, "
    "applying the validated methodology and the DrR concept to further irrigated high-value "
    "horticultural crops. Progress is measured by the career KPI K11."))

T(header=["Measure", "Content and mechanism", "Evidence"],
  rows=[
    ["Career Development Plan",
     "Jointly owned by fellow and supervisors; links the four objectives to training, teaching, "
     "publication and open-science plans; reviewed at five steering points.",
     "D1.3 (M3); reviews M6/M14/M18/M24/M27 (K11)"],
    ["Supervised two-way transfer",
     "ATK: physiological stress interpretation, validation design. ELTE: hyperspectral/UAV, "
     "scale-aware EO, predictive-modelling/ML strand. Fellow contributes physiology, phenotyping, "
     "proximal interpretation.",
     "Weekly supervision; co-supervision strand"],
    ["Transferable-skills portfolio",
     "Research integrity, open science/FAIR, generative-AI literacy, grant writing, IP and "
     "exploitation, project management, science communication — each exercised on the project's own "
     "material.",
     "CDP-tracked"],
    ["Non-academic placement",
     "Six months at AgroVIR (M25–M30): operational model evaluation, requirements engineering, "
     "interoperability assessment, exploitation planning; AgroVIR consulted from M1.",
     "D4.1, D4.2 (M30)"],
    ["Teaching and dissemination route",
     "Guest lectures, seminars and practicals at MATE and (where possible) AU Plovdiv; ≥2 "
     "peer-reviewed publications targeted, possible third methods paper.",
     "CDP; publications"],
  ], widths=[3.2, 10.2, 4.4])

P(text=(
    "By completion the fellow will have carried one scientific concept through biological definition, "
    "model development, unseen-condition validation, software implementation and industrial evaluation "
    "— a distinctive cross-sectoral profile for academic and non-academic destinations alike. The Career "
    "Development Plan makes the pathway auditable and adaptive; working-condition alignment with the "
    "European Charter for Researchers is evidenced at instrument level, and no institutional HR award "
    "is claimed."))

H(2, "2.2 Suitability and quality of the measures to maximise expected outcomes and impacts, as set out in the dissemination and exploitation plan, including communication activities")

P(text=(
    "The plan is scaled to what one experienced researcher can credibly deliver, runs as a continuous "
    "strand across the whole action, and is governed by 'as open as possible, as closed as necessary'. "
    "Five target groups are addressed through matched channels: research communities; growers and "
    "advisors in Hungary and Bulgaria; the FMIS/agri-software industry; students; and the public "
    "through the SDG 2/6/12/13 frame."))

T(header=["Strand", "Target groups", "Measures and concrete outputs", "Deliverables"],
  rows=[
    ["Dissemination",
     "Plant-physiology, crop-modelling, RS/EO and precision-agriculture communities; students",
     "≥2 peer-reviewed papers (open-access targeted, possible third); citable model card and "
     "crop-transfer specification; the prospectively measured transfer gap published, not hidden; "
     "teaching route at MATE / AU Plovdiv.",
     "D1.4 (M3); D1.5 (M14); D2.2 (M18)"],
    ["Communication",
     "Commercial vegetable growers (HU/BG); agronomic advisors; interested public",
     "Practitioner demonstrations of stress and uncertainty outputs; ≥1 public-facing article per "
     "year; Bulgarian farmer-facing demonstrations through Krumatic; conditional Szentendre Skanzen "
     "demo; no unearned numeric water-saving claims.",
     "Annual articles; demos"],
    ["Exploitation",
     "FMIS providers and agricultural-software developers; growers via them",
     "Crop-configurable DrR web MVP with validated tomato module; pre-placement dossier; AgroVIR "
     "operational/interoperability assessment; post-MSCA FMIS roadmap; MVCRI replication route; "
     "Krumatic Bulgarian uptake.",
     "D3.1, D3.2 (M24); D4.1, D4.2 (M30)"],
    ["IP management",
     "—",
     "Separate rights layers (DrR background / FIELDWISE foreground / institutional datasets / "
     "partner code); maintained inventory; dated DrR background declaration with deposit; placement "
     "agreement fixes ownership (grant-agreement default, HUN-REN ATK) and DrR access; Krumatic "
     "code-IP via pre-grant development agreement.",
     "Inventory; pre-grant agreements"],
  ], widths=[2.6, 4.2, 8.2, 2.8])

P(text=(
    "Reach and uptake are measured through owner-committed indicators. Field demonstrations target at "
    "least 15 practitioners each; public-facing articles target at least 300 views where the outlet "
    "reports analytics, otherwise a named sector outlet. Scientific "
    "dissemination targets two named events — the EGU General Assembly and the European Conference on "
    "Precision Agriculture (ECPA). The committed output set: at least two conference presentations; at "
    "least two peer-reviewed manuscripts submitted, one targeted for publication within the fellowship; "
    "a specialised training school or workshop; two institutional seminars; an industry or end-user "
    "workshop; an exploitation consultation; and a European Researchers' Night activity. The AgroVIR "
    "evaluation is completed across all six assessment dimensions (K9); repository downloads and "
    "citations of the model card and crop-transfer specification are monitored from release, not "
    "promised."))

H(2, "2.3 The magnitude and importance of the project's contribution to the expected scientific, societal and economic impacts")

P(text=(
    "The impact claim is calibrated to a single-fellow, 30-month action: FIELDWISE quantifies what it "
    "controls and contextualises the rest from published literature. The durable result is "
    "methodological — which water-stress relationships survive the transition toward field-scale "
    "monitoring, and which workflow components are genuinely reusable beyond the development crop."))

T(header=["Dimension", "Contribution (mechanism)", "Magnitude and importance"],
  rows=[
    ["Scientific",
     "A validation framework judging operational data streams against independent plant-level ground "
     "truth under unseen-year and new-field conditions; reusable-vs-crop-specific separation; citable "
     "model card and crop-transfer specification; the transfer gap stated as a result (D2.2).",
     "A portable methodology other groups adopt without opaque AI; confronts the weakness that "
     "retrospective accuracy does not survive leakage-safe or prospective testing; magnitude scales "
     "with reuse."],
    ["Agricultural & environmental",
     "Physiology- and plant-stage-aware irrigation timing; the project reports its own measured "
     "water, yield, quality and uncertainty results from prospective validation.",
     "Literature-bounded context, not a promise: published processing-tomato strategies report ≈8–30% "
     "water savings (Carucci et al. 2023; Badr et al. 2026). Serves SDG 2, 6, 12, 13; makes "
     "water–yield trade-offs visible earlier."],
    ["Technological & economic",
     "Crop-configurable DrR web MVP (validated tomato module); AgroVIR requirements, assessment and "
     "post-MSCA FMIS roadmap; Krumatic implementation and Bulgarian uptake.",
     "Uptake targets an existing operational base: the placement partner reports an FMIS in "
     "production use on >745,000 ha in eight countries (company-declared, pending written "
     "confirmation). No invented EUR/ha claims; measured operational metrics enable later economic "
     "assessment."],
    ["European & societal",
     "Hungary–Bulgaria scientific and mobility bridge; named reuse channels: industrial (AgroVIR), "
     "cross-border software (Krumatic), institutional replication (MVCRI).",
     "Potential domain, not addressable market: EU fresh vegetables ≈2.0 million ha (Eurostat 2023), "
     "≈0.7 million farms (2020 census). The realistic near-term segment is the HU/BG "
     "processing-tomato sector: in Hungary the principal Univer raw-material network integrates "
     "≈1,500 ha and nearly 100 farmers — about 70–80% of national production (industry-reported "
     "figures; citation pending, flagged); in Bulgaria ≈1,100 ha contracted by the three main "
     "processors in 2024 (WPTC tables) — an immediately relevant domain of ≈2,600 ha."],
  ], widths=[2.6, 7.0, 8.2])

P(text=(
    "The project's own claimed value is correct stress detection, warning lead time, uncertainty "
    "honesty, usability and irrigation-attention relevance; lead time is quantified prospectively with "
    "the pre-registered metric, and no fixed numerical lead-time benefit is claimed before validation. Delivery is auditable through eleven KPIs, each measuring something "
    "the project controls and each traced to a named deliverable: K1 harmonised archive coverage "
    "(D1.1); K2 predefined stress target and validation design (D1.2); K3 documented released model "
    "(D1.4); K4 crop-transfer specification (D1.5); K5 independent validation dataset (D2.1); K6 "
    "quantified transfer gap (D2.2); K7 delivered MVP (D3.1); K8 pre-placement readiness dossier "
    "(D3.2); K9 completed external evaluation (D4.1); K10 defined integration path (D4.2); K11 Career "
    "Development Plan agreed and reviewed (D1.3)."))

P(text=(
    "Magnitude within the fellowship is modest and honestly bounded — one validated crop module, one "
    "instrumented commercial field, one industrial evaluation. But every quantified magnitude is "
    "either published literature labelled as context or measured from the project's own prospective "
    "data, and the transferable core, the crop gates and the MVCRI route extend the reach beyond the "
    "24-month scientific programme."))

# ================= 3. IMPLEMENTATION =================
H(1, "3. Quality and Efficiency of the Implementation")
H(2, "3.1 Quality and effectiveness of the work plan, assessment of risks and appropriateness of the effort assigned to work packages")

P(text=(
    "FIELDWISE is a single-beneficiary 30-month European Fellowship: a 24-month research fellowship at "
    "HUN-REN ATK followed by a 6-month non-academic placement at AgroVIR. The fellow personally "
    "executes the scientific content of every work package; partners contribute supervision, data, "
    "instruments and software engineering in kind. The plan is deliberately "
    "compact — four work packages, eleven deliverables, six milestones — following the causal chain "
    "from evidence to operation; the primary model is frozen at month 3 (MS6), before the first "
    "prospective field season, so validation is genuinely out-of-sample."))

T(header=["Work package (lead)", "Months", "Content", "Deliverables"],
  rows=[
    ["WP1 Data harmonisation, stress definition & model development (HUN-REN ATK)",
     "M1–M14",
     "Harmonise the 2022–2026 MATE archive into a QC-controlled common-variable matrix; fix a "
     "physiologically meaningful stress target and a pre-registered, leakage-safe validation protocol; "
     "develop and freeze the primary model at month 3; then spectral/UAV/Sentinel scaling and the "
     "crop-transfer specification.",
     "D1.1 (M2); D1.2, D1.3 CDP, D1.4 (M3); D1.5 (M14)"],
    ["WP2 Independent commercial-farmer prospective validation (scientific lead: the fellow at the "
     "beneficiary HUN-REN ATK, under host supervision)",
     "M4–M18",
     "Written field access secured; project meteorological station and soil sensors installed; "
     "coordinated physiological, proximal, thermal, UAV and Sentinel campaigns; the frozen model "
     "applied prospectively and the transfer gap quantified with limited, transparent recalibration. "
     "The farmer hosts the validation environment; protocol, acceptance criteria, analysis and "
     "interpretation remain with the fellow.",
     "D2.1, D2.2 (M18)"],
    ["WP3 DrR web MVP, co-development & exploitation preparation (Krumatic)",
     "M1–M24",
     "Continuous co-development from month 1 with AgroVIR requirements consultation; incremental "
     "integration of validated functions; migration of the validated core; maintained "
     "background/foreground IP inventory.",
     "D3.1, D3.2 (M24)"],
    ["WP4 AgroVIR non-academic placement (AgroVIR)",
     "M25–M30",
     "Operational evaluation of the MVP in a real FMIS environment (usability, missing-data "
     "behaviour, interoperability); post-MSCA integration and exploitation roadmap.",
     "D4.1, D4.2 (M30)"],
  ], widths=[4.6, 1.6, 8.6, 3.0])

T(header=["Milestone", "Month", "Verifiable achievement criterion"],
  rows=[
    ["MS1", "M2", "Historical knowledge base and target definition complete (D1.1 accepted)."],
    ["MS6", "M3", "Primary model frozen: parameters fixed and the model card (D1.4) registered "
                  "against protocol D1.2, before the first prospective field season."],
    ["MS2", "M14", "Model and scale-transition framework ready for final prospective assessment "
                   "(D1.5)."],
    ["MS3", "M18", "Independent real-world transferability quantified with an uncertainty estimate "
                   "(D2.1, D2.2)."],
    ["MS4", "M24", "MVP ready for industrial placement evaluation — only a validated, access-tested "
                   "MVP passes into the placement (D3.1, D3.2)."],
    ["MS5", "M30", "Industrial evaluation and placement completed (D4.1, D4.2)."],
  ], widths=[1.8, 1.4, 14.6])

P(text=(
    "The single strict finish-to-start dependency is WP3→WP4: the MVP delivered at MS4 must exist "
    "before the placement evaluation begins. The remaining links are data-input dependencies, which is "
    "what lets WP3 co-develop from month 1. The critical path runs: freeze the primary model → apply "
    "it prospectively and quantify the transfer gap → migrate validated core functions into the MVP → "
    "operational evaluation. Milestone-gated go/no-go review means no package proceeds on an input "
    "that has not passed its verifiable criterion."))

P(lead="Effort.", text=(
    "The fellow's 30 person-months follow the intellectual load: WP1 9.0, WP2 6.6, WP3 6.0, WP4 5.4, "
    "plus 3.0 cross-cutting (training, dissemination, open science, management). WP1 carries the "
    "largest share because it holds the scientific core; WP2 is campaign-intensive; WP3 is spread "
    "thinly but continuously; WP4 concentrates on the placement. Partner effort is in kind and "
    "operator-declared pending each partner's written confirmation: host supervisor 2.5 PM, ELTE "
    "co-supervisor 2.0 PM, AgroVIR 3.0 PM, MATE 2.4 PM, Krumatic 1.5 PM. These contributions are "
    "deliberately structured to reduce the fellow's peak workload in the overlapping field, modelling "
    "and software period (months 4–14): ATK supervision, MATE data support, ELTE EO-campaign support "
    "and Krumatic engineering absorb work that would otherwise fall on the fellow."))

T(header=["#", "Risk (likelihood/impact)", "Mitigation"],
  rows=[
    ["R01", "Historical MATE variables differ among years (M/H)",
     "Explicit year-by-variable matrix; harmonise only defensible common variables; keep "
     "year-specific data as secondary."],
    ["R02", "Limited effective biological sample size (M/H)",
     "Repeated-measure/hierarchical structure; conservative feature reduction; parsimony, no deep "
     "learning."],
    ["R03", "Farmer site lacks monitoring infrastructure (H/M)",
     "Install project meteorological station and soil sensors before the campaign; define "
     "maintenance/data-quality responsibilities."],
    ["R04", "Primary farmer access changes or is interrupted (M/H)",
     "Written access secured before the campaign; ladder: an identified second processing-tomato "
     "producer as first-line, claim-preserving backup; the Szentendre Skanzen garden as "
     "measurement-only fallback that does not substitute for commercial-field transferability "
     "evidence — the claim is then explicitly downgraded and reported as such."],
    ["R05", "ELTE instrument access or timing constrained (M/H)",
     "Written equipment/operation schedule; prioritise key physiological dates; physiology/proximal "
     "pathway remains core if flights are missed."],
    ["R06", "Field geometry or clouds limit Sentinel-2 (H/M)",
     "Use Sentinel-2 only where valid pixels support field-scale inference; UAV/proximal carry finer "
     "scales."],
    ["R07", "Prospective performance drops (M/H)",
     "Treat transfer loss as a scientific result; quantify uncertainty; limited transparent "
     "recalibration."],
    ["R08", "Krumatic web development delayed (M/M)",
     "Incremental development from the existing prototype; scientific deliverables remain primary."],
    ["R09", "AgroVIR cannot share proprietary/customer data (L/M)",
     "Placement designed around protected/coded extracts; customer-system access is not a "
     "dependency."],
    ["R10", "MATE-derived assets carry commercial restrictions (M/H)",
     "Approved-research use only; background/foreground inventory; institutional data-rights/IP "
     "assessment before commercial exploitation."],
  ], widths=[1.2, 5.6, 11.0])

P(text=(
    "Risk is monitored, not merely catalogued: the Project Steering Group meets at M6, M14, M18, M24 "
    "and M27, issuing go/no-go verdicts on the milestones; day-to-day scientific risk is handled with "
    "the primary supervisor through the weekly cadence. Transfer loss is a scientific result to be "
    "measured, never concealed — embedded in the plan through the model-freeze-before-validation "
    "sequencing."))

IMG(GANTT_PNG,
    "Gantt chart — work-package spans, deliverables (△) and milestones (◆), months elapsed M1–M30. "
    "There are no secondments; the mobility element is the WP4 placement.")

H(2, "3.2 Quality and capacity of the host institutions and participating organisations, including hosting arrangements")

P(text=(
    "The sole beneficiary and recruiting organisation is HUN-REN ATK; full infrastructure and "
    "facilities details are in Part B-2 Section 5. Each organisation below brings a capacity the work "
    "plan genuinely depends on, with every forward-looking commitment separated from what is "
    "established."))

T(header=["Organisation (role)", "Capacity and arrangements"],
  rows=[
    ["HUN-REN ATK — beneficiary, recruiting organisation and academic host",
     "Public research organisation (HUN-REN network), Martonvásár (PIC 866599553); fellowship "
     "anchored in the Agricultural Institute. Primary supervisor Prof. Tibor Janda, Head of the Plant "
     "Physiology and Metabolomics Department (plant physiology, drought/abiotic-stress physiology, "
     "stress metabolomics); his seat acceptance and effort are stated pending the host's written "
     "hosting confirmation. Extensive drought and abiotic-stress environment; modelling collaboration "
     "with the Department of Crop Production; HUN-REN compute infrastructure. Proven MSCA host "
     "(coordinated MSCA-IF LANDRACES, 2017–2019); current participations COUSIN, AI4SoilHealth, TUdi. "
     "Workspace, integration, administrative and grants support continue through the placement; the "
     "recent HUN-REN transformation is flagged for Part A verification."],
    ["AgroVIR Kft. — non-academic placement host",
     "Hungarian FMIS SME (founded 2007, Budaörs); requirements consultation from M1; placement "
     "M25–M30 delivered through monthly structured face-to-face technical reviews and continuous "
     "remote collaboration, supervised by Miklós Maróti (Managing Director). Track record is "
     "commercial, not academic: an FMIS reported in production use on >745,000 ha in eight countries "
     "incl. Hungary and Bulgaria — declared position pending written confirmation — and no EU-funded "
     "research participation, stated openly. No proprietary or customer-data dependency "
     "(protected/coded extracts only); a written placement agreement fixes result ownership and DrR "
     "access before the placement. No PIC yet; registration is a submission-time action."],
    ["ELTE — associated partner (remote sensing and modelling)",
     "Institute of Cartography and Geoinformatics; co-supervisor Prof. András Jung; named ELTE-side "
     "team member Dr. Roland Hollós for the predictive-modelling/ML strand (representation basis "
     "pending ELTE's written confirmation); UAV operations support Dr. Zsófia Varga. GIS/RS laboratory "
     "with DJI Matrice 350 RTK (Zenmuse L1/P1/H20T), Ultris S5 hyperspectral camera, Cubert CUVIS "
     "laboratory package, ArcGIS/ENVI/MATLAB, open geospatial tools, HPC. Written instrument-access "
     "arrangement pending; risk R05 is the constrained-access fallback."],
    ["MATE — associated partner (source scientific environment)",
     "Institute of Horticultural Sciences (PIC 891269563): the 2022–2026 processing-tomato archive; "
     "soil-moisture, stomatal, SPAD, VIS-NIR/SWIR spectroradiometer, meteorological and CWSI "
     "instrumentation. Contact Dr. Sándor Takács; precision-agriculture depth from Prof. Gábor "
     "Milics. Data-use permission to be documented before the action starts (risk R10)."],
    ["Krumatic EOOD — associated partner (software)",
     "Bulgarian technology company (EIK 201511899, Sofia; manager and signatory Lyubomir Krumov): "
     "continuous DrR development, crop-configurable architecture, migration of the validated core to "
     "the web MVP, later Bulgarian farmer-facing uptake. Independence from other partners declared; "
     "code-IP terms settled in a pre-grant development agreement."],
    ["Commercial farmer site — collaborating validation environment",
     "Independent Hungarian production field, deliberately unnamed (recorded operator decision); the "
     "project installs an automated meteorological station and soil sensors and performs all "
     "specialised measurements; written field-access and treatment arrangement obtained before the "
     "campaign; ladder: second producer first (claim-preserving), Skanzen route measurement-only."],
    ["MVCRI (Maritsa VCRI, Plovdiv) — post-MSCA route",
     "Confirmed replication, local-recalibration and extension route under Bulgarian "
     "vegetable-production conditions; carries no in-action effort."],
  ], widths=[4.2, 13.6])

P(text=(
    "FIELDWISE is a European Fellowship arising from an intra-EU move (Bulgaria to Hungary); the "
    "Global Fellowship two-host requirement does not apply, and the placement is an integral "
    "extension, not an outgoing phase."))


# ---------------------------------------------------------------------------
# Gantt figure (deterministic, from Tier 4 gantt.json + display mapping)
# ---------------------------------------------------------------------------
def build_gantt() -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    g = json.loads(GANTT_JSON.read_text(encoding="utf-8-sig"))
    spans: dict[str, list[int]] = {}
    for t in g["tasks"]:
        wp = WP_DISPLAY[t["wp_id"]]
        s, e = int(t["start_month"]), int(t["end_month"])
        if wp not in spans:
            spans[wp] = [s, e]
        else:
            spans[wp][0] = min(spans[wp][0], s)
            spans[wp][1] = max(spans[wp][1], e)

    order = ["WP1", "WP2", "WP3", "WP4"]
    labels = {
        "WP1": "WP1 Data harmonisation / model",
        "WP2": "WP2 Prospective validation",
        "WP3": "WP3 DrR web MVP",
        "WP4": "WP4 AgroVIR placement",
    }
    wp_of_display_deliv = {
        "D1.1": "WP1", "D1.2": "WP1", "D1.3": "WP1", "D1.4": "WP1", "D1.5": "WP1",
        "D2.1": "WP2", "D2.2": "WP2", "D3.1": "WP3", "D3.2": "WP3",
        "D4.1": "WP4", "D4.2": "WP4",
    }

    fig, ax = plt.subplots(figsize=(9.6, 2.9), dpi=200)
    y = {wp: len(order) - i for i, wp in enumerate(order)}
    for wp in order:
        s, e = spans[wp]
        ax.barh(y[wp], e - s + 1, left=s - 0.5, height=0.5,
                color="#4472a8", edgecolor="#2c4a70", zorder=2)

    by_pos: dict[tuple, list[str]] = {}
    for internal, (disp, month) in DELIVERABLE_DISPLAY.items():
        by_pos.setdefault((wp_of_display_deliv[disp], month), []).append(disp)
    last_label_month: dict[str, int] = {}
    for (wp, month), ids in sorted(by_pos.items()):
        ax.plot(month, y[wp] + 0.42, marker="^", color="#1a1a1a", markersize=6,
                linestyle="none", zorder=3)
        # stagger labels when two marker groups sit within 2 months of each other
        close = wp in last_label_month and month - last_label_month[wp] <= 2
        dy = 0.78 if close else 0.52
        ax.annotate(",".join(sorted(ids)), (month, y[wp] + dy), ha="center",
                    fontsize=6.2, color="#1a1a1a")
        last_label_month[wp] = month
    for ms, month, wp in MILESTONES:
        ax.plot(month, y[wp] - 0.42, marker="D", color="#a83232", markersize=5,
                linestyle="none", zorder=3)
        ax.annotate(ms, (month, y[wp] - 0.72), ha="center", fontsize=6.2,
                    color="#a83232")

    ax.set_yticks([y[wp] for wp in order])
    ax.set_yticklabels([labels[wp] for wp in order], fontsize=7.5)
    ax.set_xlim(0.5, 30.5)
    ax.set_ylim(0.2, len(order) + 0.9)
    ax.set_xticks(range(1, 31))
    ax.tick_params(axis="x", labelsize=6.5)
    ax.set_xlabel("Month (elapsed, M1–M30)", fontsize=7.5)
    ax.grid(axis="x", color="#d9d9d9", linewidth=0.5, zorder=1)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(GANTT_PNG, metadata={})
    plt.close(fig)


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------
def render() -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, attr, Cm(1.5))

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(4)

    def add_heading(level: int, text: str) -> None:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Times New Roman"
        if level == 0:
            run.font.size = Pt(14)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif level == 1:
            run.font.size = Pt(12.5)
            p.paragraph_format.space_before = Pt(8)
        else:
            run.font.size = Pt(11)
            p.paragraph_format.space_before = Pt(6)

    def add_para(lead: str, text: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if lead:
            r = p.add_run(lead + " ")
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    def add_table(header: list, rows: list, widths: list | None) -> None:
        tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl.style = "Table Grid"
        data = [header] + rows
        for i, row in enumerate(data):
            for j, cell_text in enumerate(row):
                cell = tbl.rows[i].cells[j]
                cell.text = ""
                para = cell.paragraphs[0]
                run = para.add_run(cell_text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                run.bold = i == 0
                para.paragraph_format.space_after = Pt(1)
        if widths:
            for j, w in enumerate(widths):
                for i in range(len(data)):
                    tbl.rows[i].cells[j].width = Cm(w)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    def add_image(path: str, caption: str) -> None:
        doc.add_picture(path, width=Cm(17.8))
        p = doc.add_paragraph()
        r = p.add_run(caption)
        r.italic = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(9)

    for b in C["blocks"]:
        if b["kind"] == "h":
            add_heading(b["level"], b["text"])
        elif b["kind"] == "p":
            add_para(b["lead"], b["text"])
        elif b["kind"] == "t":
            add_table(b["header"], b["rows"], b["widths"])
        elif b["kind"] == "img":
            add_image(b["path"], b["caption"])

    doc.save(OUT_DOCX)


# ---------------------------------------------------------------------------
# Self-checks (mechanical half of the plan §6.4 pre-freeze checklist)
# ---------------------------------------------------------------------------
def full_text() -> str:
    parts = []
    for b in C["blocks"]:
        if b["kind"] == "h":
            parts.append(b["text"])
        elif b["kind"] == "p":
            parts.append((b["lead"] + " " + b["text"]).strip())
        elif b["kind"] == "t":
            for row in [b["header"]] + b["rows"]:
                parts.extend(row)
        elif b["kind"] == "img":
            parts.append(b["caption"])
    return "\n".join(parts)


def self_check() -> list[str]:
    text = full_text()
    errors: list[str] = []

    def forbid(needle: str, why: str) -> None:
        if needle.lower() in text.lower():
            errors.append(f"forbidden phrase present ({why}): {needle!r}")

    def require(needle: str, why: str) -> None:
        if needle not in text:
            errors.append(f"required content missing ({why}): {needle!r}")

    forbid("single remaining", "T1 wording replacement")
    forbid("mainly remote", "T7 placement reframe")
    forbid("HUN-REN-led", "flag 3: modelling leadership is ELTE-side")
    forbid("[A-", "anchor markup must be stripped")
    forbid("¶", "anchor markup must be stripped")
    forbid("WP5", "display mapping: only WP1–WP4 exist")
    forbid("proves water savings", "T5 decision-value discipline")

    if re.search(r"\bT\d+\.\d+\b", text):
        errors.append("task ids present — condensed form must describe tasks verbally")
    for m in re.finditer(r"\bD(\d)\.(\d)\b", text):
        disp = m.group(0)
        allowed = {d for d, _ in DELIVERABLE_DISPLAY.values()}
        if disp not in allowed:
            errors.append(f"non-display deliverable id used: {disp}")
    for disp, _ in DELIVERABLE_DISPLAY.values():
        require(disp, "every display deliverable id must appear")
    for m in re.finditer(r"\bMS\d+\b", text):
        if m.group(0) not in {ms for ms, _, _ in MILESTONES}:
            errors.append(f"unknown milestone id: {m.group(0)}")
    for k in range(1, 12):
        require(f"K{k}", "T3: all eleven KPIs named")

    require("principal methodological competence gap", "T1 replacement wording")
    require("citation pending", "A-10 §13.8 flag on the Univer segment figures")
    require("EGU General Assembly", "A-9/A-12 named events")
    require("European Conference on Precision Agriculture", "A-9/A-12 named events")
    require("745,000", "AgroVIR footprint (with pending-confirmation flag)")
    require("pending written confirmation", "flag 5: Assumed content stays visible")
    require("scientific lead: the fellow at the beneficiary HUN-REN ATK", "A-6 WP2 lead")
    require("Prof. Tibor Janda", "run-03 supervisor spine")
    require("month 2", "A-1: D1.1 due month")
    require("independently measured physiological water-stress target", "T2 model sentence")

    pend = len(re.findall(r"pending", text, re.IGNORECASE))
    if pend < 5:
        errors.append(f"only {pend} 'pending' transparency flags; expected >=5 (flag 5)")

    words = len(text.split())
    if words > 5450:
        errors.append(f"word budget exceeded: {words} words (cap 5450 for the 10-page limit; the "
                      "run-02 master's measured density is 5069 words = 9 pages)")
    print(f"[info] condensed content: {words} words")

    heads = [b["text"] for b in C["blocks"] if b["kind"] == "h"]
    for hid in ["1. Excellence", "1.1 ", "1.2 ", "1.3 ", "1.4 ", "2. Impact",
                "2.1 ", "2.2 ", "2.3 ", "3. Quality", "3.1 ", "3.2 "]:
        if not any(h.startswith(hid) for h in heads):
            errors.append(f"template section missing: {hid.strip()}")
    return errors


def main() -> int:
    build_gantt()
    errors = self_check()
    if errors:
        for e in errors:
            print(f"[FAIL] {e}")
        return 1
    render()
    print(f"[ok] gantt figure: {GANTT_PNG.relative_to(REPO)}")
    print(f"[ok] rendered: {OUT_DOCX.relative_to(REPO)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
