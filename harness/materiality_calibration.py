"""
E3.1 — Materiality-classifier calibration runner (turnkey, native).

The §10.5 materiality classifier (``harness.materiality.classify_materiality``) is
what separates a *material claim* (a checkable fact §10.5 requires be attributed
or flagged) from framing/transition prose.  Ticket E3 shipped it with **recall
computable but precision withheld** — with zero labelled negatives the confusion
matrix's precision is a spurious ``tp/(tp+0)=1.0`` (see
``MaterialityCalibration``).  E3.1 closes that: a human labels real negatives (so
precision becomes real) and vets a set of out-of-ledger *escapes* (so the
recall that matters — catching a claim that fabricates without a ledger record —
is measured directly, not just recall over the auto-positives).

This runner is the bridge from a human-friendly **labelling sheet** to the
harness's own calibration primitives.  It:

  1. reads the labelled sheet (``materiality_gold_LABEL_ME.xlsx`` — columns, a
     TRUE/FALSE dropdown; a ``.csv`` twin, comma or semicolon, also works),
  2. writes the canonical JSONL materiality set the module consumes
     (``harness.materiality.write_materiality_template`` +
     ``materiality_set_hash``),
  3. runs the **real** classifier once per row
     (``classify_materiality(judge, text, property_key=…)``) and scores it with
     the **real** confusion matrix (``harness.calibration.confusion_from_labels``),
  4. adds the E3.1-specific *escape-recall* slice over the seed rows,
  5. applies the fail-safe graduation rule (mirrors
     ``harness.calibration.graduation_for``: undefined ⇒ not graduated; a repin
     or an unreviewed label set ⇒ advisory), and
  6. writes a pre-filled **decision-log draft** (house schema; narrative fields
     left as explicit placeholders — the run reports numbers, a human authors
     the rationale, §13.3).

Nothing here is a runtime gate.  It measures the classifier the pipeline runs;
it never blocks a run (harness/HARNESS.md).

Workflow
--------
    1. Label ``harness/labeling/materiality_gold_LABEL_ME.xlsx`` — column B
       (material) is pre-filled with AI SUGGESTIONS; read every row and correct
       B where you disagree (column C keeps the suggestion, so your changes are
       visible).  Full rubric: ``harness/labeling/LABELLING_GUIDE.md``.
    2. Attest the review (required before this may ever gate):
           export HARNESS_LABELS_REVIEWED_BY="Your Name"
    3. Run (local judge over the OpenAI-compatible transport; cost $0):
           HARNESS_LABELS_REVIEWED_BY="Your Name" \
           HARNESS_JUDGE_MODEL=<local-model> HARNESS_JUDGE_VERSION=<tag> \
           ORCHESTRATOR_TRANSPORT_PRESET=openai_compatible \
           ORCHESTRATOR_TRANSPORT_ENDPOINT=http://localhost:11434/v1 \
           python -m harness.materiality_calibration

Point at a specific sheet with ``HARNESS_MATERIALITY_GOLD=/abs/path/to/sheet``.

Constitutional authority:
    Subordinate to CLAUDE.md.  Out-of-band QA substrate; reads a harness-owned
    gold sheet and writes only harness-owned artifacts.  Never a runtime gate
    (§17.6.2).  The judge's verdicts are Inferred (§12.2); the gold labels are
    human ground truth (the E1.5 independence rule).
"""
from __future__ import annotations

import csv
import datetime
import json
import os
import pathlib
from dataclasses import dataclass

# NOTE: every ``harness.*`` import lives INSIDE a function, so this module imports
# with the stdlib alone (its sheet-parsing and metric helpers are unit-testable
# without the judge/transport stack).

# ─────────────────────────────────────────────────────────────────────────────
# Operator policy — the two bars E3.1 measures. These are policy, not physics:
# pick them from the cost of a blessed escape vs an over-flag on your prose.
# (harness.calibration.DEFAULT_GRADUATION_THRESHOLD is 0.90/0.80 — a strictly
#  more conservative starting point; raise MIN_PRECISION to it if you prefer.)
MIN_PRECISION = 0.85          # of spans the judge calls material, the share truly material
MIN_RECALL_ESCAPES = 0.80     # of seeded out-of-ledger material claims, the share caught
# ─────────────────────────────────────────────────────────────────────────────

_TRUE = {"true", "t", "1", "yes", "y"}
_FALSE = {"false", "f", "0", "no", "n"}

# Where the labelled sheet may live (first hit wins). HARNESS_MATERIALITY_GOLD overrides.
_CANDIDATE_DIRS = [pathlib.Path("harness/labeling"), pathlib.Path("harness/materiality_sets")]
_CANDIDATE_NAMES = [
    "materiality_gold_LABEL_ME.xlsx", "materiality_gold_LABEL_ME.csv",
    "materiality_gold.xlsx", "materiality_gold.csv",
]


@dataclass
class Row:
    """One labelled span read from the sheet (kept independent of the harness)."""
    example_id: str
    material: bool
    ai_suggested: bool | None
    kind: str            # "seed" (out-of-ledger escape) | "negative" (real prose)
    section: str
    in_ledger: str
    note: str
    text: str
    pred: bool | None = None   # filled by the classifier during a run

    @property
    def is_seed(self) -> bool:
        return self.kind == "seed"


# ─────────────────────────────────────────────────────────────────────────────
# Sheet loading (pure: xlsx or csv; fail-closed on blank labels)
# ─────────────────────────────────────────────────────────────────────────────
def find_gold() -> pathlib.Path:
    override = os.environ.get("HARNESS_MATERIALITY_GOLD")
    if override:
        p = pathlib.Path(override)
        if not p.exists():
            raise SystemExit(f"HARNESS_MATERIALITY_GOLD points at a missing file: {p}")
        return p
    for d in _CANDIDATE_DIRS:
        for n in _CANDIDATE_NAMES:
            p = d / n
            if p.exists():
                return p
    raise SystemExit(
        "could not find the gold sheet. Looked in "
        + ", ".join(str(d) for d in _CANDIDATE_DIRS)
        + " — set HARNESS_MATERIALITY_GOLD=/abs/path/to/sheet"
    )


def _coerce_bool(v) -> bool | None:
    if v is None:
        return None
    s = str(v).strip().lower()
    if s in _TRUE:
        return True
    if s in _FALSE:
        return False
    return None


def _rows_from_xlsx(path: pathlib.Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as e:  # pragma: no cover - environment guard
        raise SystemExit(
            "reading .xlsx needs openpyxl (`pip install openpyxl`), or save the sheet as CSV."
        ) from e
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["gold"] if "gold" in wb.sheetnames else wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip() if h is not None else "" for h in next(it)]
    except StopIteration:
        return []
    out = []
    for raw in it:
        if raw is None or all(c is None for c in raw):
            continue
        out.append({header[i]: (raw[i] if i < len(raw) else None) for i in range(len(header))})
    return out


def _rows_from_csv(path: pathlib.Path) -> list[dict]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        first = f.readline()
        delim = ";" if first.count(";") > first.count(",") else ","
        f.seek(0)
        return list(csv.DictReader(f, delimiter=delim))


def read_sheet(path: pathlib.Path) -> list[Row]:
    """Load the labelled sheet; skip any row whose `material` is blank (fail-closed)."""
    raw = _rows_from_xlsx(path) if path.suffix.lower() in (".xlsx", ".xlsm") else _rows_from_csv(path)
    rows: list[Row] = []
    for r in raw:
        m = _coerce_bool(r.get("material"))
        if m is None:
            continue  # unlabelled → not counted (same posture as load_materiality_set)
        rows.append(
            Row(
                example_id=str(r.get("id") or "").strip(),
                material=m,
                ai_suggested=_coerce_bool(r.get("ai_suggested")),
                kind=str(r.get("kind") or "").strip(),
                section=str(r.get("section") or "").strip(),
                in_ledger=str(r.get("in_ledger") or "").strip(),
                note=str(r.get("note") or "").strip(),
                text=str(r.get("text") or ""),
            )
        )
    return rows


# ─────────────────────────────────────────────────────────────────────────────
# E3.1 metrics + graduation (pure)
# ─────────────────────────────────────────────────────────────────────────────
def escape_recall(rows: list[Row]) -> float | None:
    """Recall over the seed rows (all material by construction): the share caught.

    This is the dangerous-miss metric E3.1 exists for — a *seed* is a material
    claim with NO covering ledger record, so a seed the classifier calls
    non-material is an escape it would wave through.
    """
    seeds = [r for r in rows if r.is_seed]
    if not seeds:
        return None
    caught = sum(1 for r in seeds if r.pred)
    return caught / len(seeds)


def disagreements(rows: list[Row]) -> tuple[list[Row], list[Row], list[Row]]:
    """(missed escapes, missed embedded claims, false alarms) — human-readable errors."""
    missed_escapes = [r for r in rows if r.is_seed and not r.pred]
    missed_embedded = [r for r in rows if (not r.is_seed) and r.material and not r.pred]
    false_alarms = [r for r in rows if (not r.material) and r.pred]
    return missed_escapes, missed_embedded, false_alarms


def decide(precision, recall_escapes, reviewed: bool) -> tuple[str, str]:
    """Fail-safe graduation, mirroring ``harness.calibration.graduation_for``.

    Gating is permitted only when BOTH bars are cleared AND a human has attested
    the labels.  Undefined precision/recall (``None``) never clears the bar; an
    unreviewed (AI-suggested) label set is pinned to advisory regardless of the
    numbers — the harness may not gate on labels a human has not confirmed.
    """
    metrics_clear = (
        precision is not None and precision >= MIN_PRECISION
        and recall_escapes is not None and recall_escapes >= MIN_RECALL_ESCAPES
    )
    decision = "gating_permitted" if (metrics_clear and reviewed) else "advisory"
    labels = (
        f"human-reviewed ({reviewed})" if isinstance(reviewed, str) and reviewed
        else "AI-SUGGESTED — UNCONFIRMED (independence not established; set HARNESS_LABELS_REVIEWED_BY)"
    )
    return decision, labels


# ─────────────────────────────────────────────────────────────────────────────
# The run (touches the harness: judge, classifier, native confusion matrix)
# ─────────────────────────────────────────────────────────────────────────────
def _build_judge():
    """Return (config, judge): the pinned, non-drafter judge over the OAI transport (E1)."""
    from harness.judge import resolve_judge_config, Judge
    from harness.provenance import ProvenanceLog
    cfg = resolve_judge_config()
    prov_path = _out_dir() / "provenance_materiality_calibration.jsonl"
    judge = Judge(cfg, provenance_log=ProvenanceLog(prov_path))
    return cfg, judge


def _to_examples(rows: list[Row]):
    """Build the canonical MaterialityExample list (origin=prose; seed-ness kept in note)."""
    from harness.materiality import MaterialityExample, ORIGIN_PROSE
    return [
        MaterialityExample(
            example_id=r.example_id,
            text=r.text,
            material=r.material,
            origin=ORIGIN_PROSE,
            section_id=r.section or None,
            note=f"{r.kind}: {r.note}".strip(": ").strip(),
        )
        for r in rows
    ]


_GOLD_PATH: pathlib.Path | None = None


def _out_dir() -> pathlib.Path:
    return (_GOLD_PATH.parent if _GOLD_PATH is not None else pathlib.Path("harness/labeling"))


def main() -> None:
    global _GOLD_PATH
    _GOLD_PATH = find_gold()
    rows = read_sheet(_GOLD_PATH)
    if not rows:
        raise SystemExit(
            f"no labelled rows in {_GOLD_PATH} — fill the `material` column (TRUE/FALSE) and save"
        )

    reviewer = (os.environ.get("HARNESS_LABELS_REVIEWED_BY") or "").strip()
    overrides = sum(
        1 for r in rows if r.ai_suggested is not None and r.material != r.ai_suggested
    )

    # --- native pieces ------------------------------------------------------- #
    from harness.materiality import (
        MaterialityCalibration, classify_materiality, write_materiality_template,
        materiality_set_hash,
    )
    from harness.calibration import confusion_from_labels

    cfg, judge = _build_judge()
    examples = _to_examples(rows)

    # Persist the canonical JSONL set the module consumes, and hash it.
    generated = _out_dir() / "materiality_gold.generated.jsonl"
    write_materiality_template(examples, generated)
    set_hash = materiality_set_hash(examples)

    # One real classification pass — the exact path the pipeline runs. (This is
    # calibrate_materiality's classifier + confusion inline, so we also keep the
    # per-row predictions the escape-recall slice and the disagreement view need.)
    n = int(os.environ.get("HARNESS_MATERIALITY_N", "1"))
    for row, ex in zip(rows, examples):
        verdict = classify_materiality(judge, ex.text, property_key=ex.example_id, n=n)
        if verdict.passed is None:
            raise SystemExit(
                f"example {ex.example_id!r}: classifier returned no boolean verdict; cannot calibrate."
            )
        row.pred = bool(verdict.passed)

    truth = [r.material for r in rows]
    pred = [bool(r.pred) for r in rows]
    confusion = confusion_from_labels(truth, pred)             # positive class = material
    labeled_negatives = sum(1 for t in truth if not t)
    calib = MaterialityCalibration(
        judge_model=cfg.model,
        judge_version=cfg.version,
        recall=confusion.recall,                              # material recall over ALL rows
        precision=(confusion.precision if labeled_negatives > 0 else None),  # None-guard (E3)
        confusion=confusion,
        n_examples=len(rows),
        labeled_negatives=labeled_negatives,
        set_hash=set_hash,
        timestamp=datetime.datetime.now(datetime.timezone.utc).isoformat(),
    )
    recall_escapes = escape_recall(rows)
    decision, labels_provenance = decide(calib.precision, recall_escapes, reviewer or False)

    # --- human-readable view ------------------------------------------------- #
    seeds = [r for r in rows if r.is_seed]
    pctf = lambda v: "n/a" if v is None else f"{v:.0%}"
    print(f"\ngold sheet: {_GOLD_PATH}")
    print(f"generated set: {generated}  ({set_hash[:23]}…)")
    print(f"judge: {cfg.model}@{cfg.version}")
    print(f"labels: {labels_provenance}")
    if not reviewer:
        print("  ⚠ no HARNESS_LABELS_REVIEWED_BY — decision pinned to 'advisory' regardless of metrics.")
    print(f"  human overrides of AI suggestion: {overrides}/{len(rows)}"
          + ("" if overrides else "  (labels identical to AI suggestions — did you review them?)"))
    print(f"rows: {len(rows)}  (seeds {len(seeds)}, negatives {len(rows) - len(seeds)}, "
          f"labelled-negatives {labeled_negatives})")
    print(f"precision {pctf(calib.precision)}  |  material-recall(all) {pctf(calib.recall)}  |  "
          f"escape-recall(seeds) {pctf(recall_escapes)}  |  decision: {decision}\n")

    missed_escapes, missed_embedded, false_alarms = disagreements(rows)

    def show(title, rs):
        if rs:
            print(title)
            for r in rs:
                print(f"   {r.example_id}: {r.text[:88]}")
            print()

    show("MISSED escapes  (out-of-ledger material claim called non-material — the dangerous error):", missed_escapes)
    show("MISSED embedded claims  (a fact hidden in framing, waved through):", missed_embedded)
    show("FALSE alarms  (framing wrongly flagged as a material claim):", false_alarms)
    if not (missed_escapes or missed_embedded or false_alarms):
        print("no disagreements — the classifier matched every human label.\n")

    # --- decision-log DRAFT (house schema; narrative left to a human, §13.3) -- #
    stamp = datetime.date.today().isoformat()
    iso = datetime.datetime.now(datetime.timezone.utc).isoformat()
    draft = {
        "decision_id": f"e3.1_materiality_calibration_{stamp}",
        "decision_type": "eval_harness_e3.1_materiality_calibration_run",
        "invoking_agent": "harness/materiality_calibration.py (E3.1)",
        "phase_context": (
            "Out-of-band QA/CI track (harness/). First human-labelled calibration of the "
            "§10.5 materiality classifier: labelled negatives make precision real (E3 shipped "
            "it None-guarded) and seeded out-of-ledger escapes measure the miss rate that matters. "
            "Never a runtime gate (§17.6.2); advisory by construction."
        ),
        "run_id_reference": "harness-branch (offline; local judge over the OpenAI-compatible transport)",
        "authority": [
            "harness_plan/tickets_eval_harness.md E3.1 (Materiality calibration completeness)",
            "harness/materiality.py (classify_materiality + MaterialityCalibration None-guard)",
            "harness/calibration.py graduation_for (fail-safe advisory→gating cascade)",
            "CLAUDE.md §10.5 (material claims traceable; unattributed flagged, not asserted)",
            "E1.5 independence rule (human ground-truth labels; no AI labelling for an AI judge)",
        ],
        "judge": {"model": cfg.model, "version": cfg.version},
        "gold_set": {
            "source_sheet": str(_GOLD_PATH),
            "generated_jsonl": str(generated),
            "set_hash": set_hash,
            "labelled_rows": len(rows),
            "seeds": len(seeds),
            "negatives": len(rows) - len(seeds),
            "labeled_negatives": labeled_negatives,
            "labels": labels_provenance,
            "human_overrides_of_ai_suggestion": overrides,
        },
        "calibration_result": {
            **calib.to_dict(),
            "recall_escapes": recall_escapes,
            "disagreements": {
                "missed_escapes": [r.example_id for r in missed_escapes],
                "missed_embedded": [r.example_id for r in missed_embedded],
                "false_alarms": [r.example_id for r in false_alarms],
            },
        },
        "threshold": {
            "min_precision": MIN_PRECISION,
            "min_recall_escapes": MIN_RECALL_ESCAPES,
            "note": "operator policy — tune; harness default is 0.90/0.80.",
        },
        "decision": decision,
        "rationale": "<AUTO-DRAFT — human to complete: why advisory-stays or promote-to-gating, given the metrics and the disagreements above>",
        "constitutional_check": "<AUTO-DRAFT — human to confirm: E3.1 evaluates no gate (§17.6.2); advisory=True by construction; judge verdicts Inferred (§12.2); labels are human ground truth (E1.5 independence)>",
        "open_items": [
            (None if reviewer else
             "Labels are AI-suggested and UNCONFIRMED — a human must review the sheet and set "
             "HARNESS_LABELS_REVIEWED_BY before this calibration may gate."),
            (None if labeled_negatives else
             "No labelled negatives — precision is withheld (None); label negative rows FALSE to unlock it."),
            "Decision is keyed to this exact judge model+version; a repin re-opens advisory (mirrors E1.5).",
        ],
        "resolution_status": "draft",
        "timestamp": iso,
    }
    draft["open_items"] = [x for x in draft["open_items"] if x]
    out = _out_dir() / f"e3.1-materiality-calibration_DRAFT_{stamp}.json"
    out.write_text(json.dumps(draft, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {out}")
    print("→ review it, complete the two <AUTO-DRAFT> fields, then move it into "
          "docs/tier4_orchestration_state/decision_log/ as e3.1-materiality-calibration_"
          f"{stamp}.json (resolution_status: resolved).")


if __name__ == "__main__":
    main()
